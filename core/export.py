"""Excel export: one workbook per meeting, a tab per output, IDs linking across tabs."""

from __future__ import annotations

import io
import re
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from core import config
from core.formatting import email_to_text
from core.pipeline import Pipeline

FONT = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
FLAG_FILL = PatternFill("solid", fgColor="FFF2CC")
MISSING_FONT = Font(name=FONT, size=10, color="C00000", italic=True)
BODY = Font(name=FONT, size=10)
WRAP_TOP = Alignment(wrap_text=True, vertical="top")

SHEET_NAMES = ["Summary", "Requirements", "Decisions", "Epics", "Stories", "Tasks", "RAID",
               "Open Questions", "Follow-up Email"]


def export_filename(meeting: dict) -> str:
    slug = re.sub(r"[^A-Za-z0-9]+", "-", meeting["title"]).strip("-")[:50] or "meeting"
    return f"{meeting['date']}_{slug}.xlsx"


def build_workbook(pipe: Pipeline, meeting_id: int) -> bytes:
    """Build the workbook from each stage's approved output, or its latest draft."""
    db = pipe.storage
    meeting = db.get_meeting(meeting_id)
    project = db.get_project(meeting["project_id"])
    states = pipe.status(meeting_id)
    outputs = {s: (st.run["output"] if st.run else None) for s, st in states.items()}
    dor = pipe.dor_results(meeting_id)

    wb = Workbook()
    wb.remove(wb.active)
    sheets = {name: wb.create_sheet(name) for name in SHEET_NAMES}

    extract = outputs["extract"] or {"requirements": [], "decisions": [], "open_questions": []}
    _table(sheets["Requirements"],
           [("ID", 10), ("Requirement", 60), ("Type", 16), ("Speaker", 18), ("Time", 8), ("Source quote", 60)],
           [[r["id"], r["statement"], r["type"], r["speaker"], r["timestamp"], r["source_quote"]]
            for r in extract["requirements"]])
    _table(sheets["Decisions"],
           [("ID", 10), ("Decision", 60), ("Decided by", 18), ("Time", 8), ("Source quote", 60)],
           [[d["id"], d["statement"], d["decided_by"], d["timestamp"], d["source_quote"]]
            for d in extract["decisions"]])
    _table(sheets["Open Questions"],
           [("ID", 8), ("Question", 60), ("Raised by", 18), ("Owner", 18), ("Linked to", 20), ("Resolved", 10)],
           [[q["id"], q["question"], q["raised_by"], q["owner"], ", ".join(q["linked_ids"]),
             "Yes" if q["resolved"] else "No"] for q in extract["open_questions"]],
           missing_cols={3})

    epics = (outputs["epics"] or {"epics": []})["epics"]
    _table(sheets["Epics"], [("ID", 10), ("Epic", 34), ("Description", 70), ("Requirements", 30)],
           [[e["id"], e["title"], e["description"], ", ".join(e["requirement_ids"])] for e in epics])

    stories = (outputs["stories"] or {"stories": []})["stories"]
    story_rows, flagged = [], set()
    for i, s in enumerate(stories):
        result = dor.get(s["id"])
        ok = result is None or result.passed
        if not ok:
            flagged.add(i)
        ac_text = "\n".join(
            f"{n}. {'[Negative] ' if ac['negative'] else ''}Given {ac['given']}, when {ac['when']}, then {ac['then']}"
            for n, ac in enumerate(s["acceptance_criteria"], start=1))
        story_rows.append([
            s["id"], s["epic_id"], f"As a {s['as_a']}, I want {s['i_want']}, so that {s['so_that']}.",
            ac_text, s["priority"], s["points"], ", ".join(s["requirement_ids"]),
            "Pass" if ok else "Needs work", "\n".join(result.failures) if result else "",
        ])
    _table(sheets["Stories"],
           [("ID", 10), ("Epic", 10), ("User story", 50), ("Acceptance criteria", 70), ("Priority", 10),
            ("Points", 8), ("Requirements", 18), ("DoR", 12), ("DoR issues", 50)],
           story_rows, flagged_rows=flagged)

    tasks_out = outputs["tasks"] or {"tasks": [], "missing_info": []}
    epic_of = {s["id"]: s["epic_id"] for s in stories}
    questions = {}
    for m in tasks_out["missing_info"]:
        questions.setdefault(m["item_id"], []).append(m["question"])
    _table(sheets["Tasks"],
           [("ID", 10), ("Story", 10), ("Epic", 10), ("Task", 60), ("Owner", 18), ("Depends on", 18),
            ("System", 26), ("Open questions", 50)],
           [[t["id"], t["story_id"], epic_of.get(t["story_id"], ""), t["description"], t["owner"],
             ", ".join(t["dependency_ids"]), t["system"], "\n".join(questions.get(t["id"], []))]
            for t in tasks_out["tasks"]],
           missing_cols={4}, flagged_rows={i for i, t in enumerate(tasks_out["tasks"]) if t["id"] in questions})

    raid_out = outputs["raid_email"]
    _table(sheets["RAID"],
           [("ID", 9), ("Type", 13), ("Description", 55), ("Owner", 18), ("Status", 12), ("Impact", 10),
            ("Likelihood", 11), ("Mitigation / next step", 50), ("Due date", 12), ("Linked to", 20)],
           [[r["id"], r["type"], r["description"], r["owner"], r["status"], r["impact"], r["likelihood"],
             r["mitigation"], r["due_date"], ", ".join(r["linked_ids"])] for r in (raid_out or {"raid": []})["raid"]],
           missing_cols={3})

    _email(sheets["Follow-up Email"], raid_out["email"] if raid_out else None)
    _summary(sheets["Summary"], project, meeting, states, db.meeting_cost(meeting_id))

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _table(ws: Worksheet, columns: list[tuple[str, int]], rows: list[list], *,
           flagged_rows: set[int] = frozenset(), missing_cols: set[int] = frozenset()) -> None:
    """Header row, frozen and filterable; wrapped body; flagged rows shaded; blank owners shown as TBD."""
    for c, (title, width) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=c, value=title)
        cell.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = width
    for r, row in enumerate(rows, start=2):
        for c, value in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c)
            if value in (None, "") and (c - 1) in missing_cols:
                cell.value, cell.font = "TBD", MISSING_FONT
            else:
                cell.value, cell.font = value, BODY
            cell.alignment = WRAP_TOP
            if (r - 2) in flagged_rows:
                cell.fill = FLAG_FILL
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{max(len(rows) + 1, 1)}"


def _email(ws: Worksheet, email: dict | None) -> None:
    ws.column_dimensions["A"].width = 100
    lines = email_to_text(email).splitlines() if email else ["Not generated yet."]
    for r, line in enumerate(lines, start=1):
        cell = ws.cell(row=r, column=1, value=line)
        cell.font = Font(name=FONT, size=10, bold=r == 1)
        cell.alignment = Alignment(wrap_text=True)


def _summary(ws: Worksheet, project: dict, meeting: dict, states: dict, cost: float) -> None:
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 60
    ws["A1"] = "Meeting-to-Backlog export"
    ws["A1"].font = Font(name=FONT, size=14, bold=True)

    models = sorted({st.run["model"] for st in states.values() if st.run and st.run["model"]})
    info = [
        ("Project", project["name"]),
        ("Meeting", meeting["title"]),
        ("Meeting date", meeting["date"]),
        ("Exported", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")),
        ("Model", ", ".join(models) or "—"),
        ("Model cost (USD)", round(cost, 2)),
    ]
    # Counts are formulas over the other tabs, so they stay right if rows are edited in Excel.
    counts = [
        ("Requirements", "=COUNTA(Requirements!A:A)-1"),
        ("Decisions", "=COUNTA(Decisions!A:A)-1"),
        ("Open questions", "=COUNTA('Open Questions'!A:A)-1"),
        ("Epics", "=COUNTA(Epics!A:A)-1"),
        ("Stories", "=COUNTA(Stories!A:A)-1"),
        ("Story points", "=SUM(Stories!F:F)"),
        ("Tasks", "=COUNTA(Tasks!A:A)-1"),
        ("Tasks without an owner", '=COUNTIF(Tasks!E:E,"TBD")'),
        ("RAID items", "=COUNTA(RAID!A:A)-1"),
        ("Definition of Ready pass rate", '=IFERROR(COUNTIF(Stories!H:H,"Pass")/(COUNTA(Stories!A:A)-1),0)'),
    ]
    row = 3
    for label, value in info + [("", None)] + counts:
        if label:
            ws.cell(row=row, column=1, value=label).font = Font(name=FONT, size=10, bold=True)
            cell = ws.cell(row=row, column=2, value=value)
            cell.font = BODY
            cell.alignment = Alignment(horizontal="left")
            if label == "Model cost (USD)":
                cell.number_format = "$#,##0.00"
            elif label == "Definition of Ready pass rate":
                cell.number_format = "0%"
        row += 1

    row += 1
    ws.cell(row=row, column=1, value="Stage").font = Font(name=FONT, size=10, bold=True)
    ws.cell(row=row, column=2, value="Status").font = Font(name=FONT, size=10, bold=True)
    for stage, st in states.items():
        row += 1
        ws.cell(row=row, column=1, value=config.STAGE_LABELS[stage]).font = BODY
        status = {"approved": "Approved", "draft": "Draft (not approved)", "stale": "Out of date",
                  "not_started": "Not generated", "blocked": "Not generated"}[st.status]
        ws.cell(row=row, column=2, value=status).font = BODY
    row += 2
    note = ws.cell(row=row, column=1, value="Rows shaded yellow need attention: stories that fail the Definition "
                                            "of Ready, and tasks with open questions. 'TBD' marks a missing owner.")
    note.font = Font(name=FONT, size=9, italic=True)
