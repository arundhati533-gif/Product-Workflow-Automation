import io
import re

import pytest
from openpyxl import load_workbook

from core import config
from core.demo import ReplayLLM, load_recordings
from core.export import SHEET_NAMES, build_workbook, export_filename
from core.parsing import parse_transcript
from core.pipeline import Pipeline
from core.samples import load_sample_project, sample_transcript_path
from core.storage import Storage


@pytest.fixture
def exported():
    storage = Storage(":memory:")
    pid = load_sample_project(storage)
    path = sample_transcript_path("01-kickoff.vtt")
    t = parse_transcript(filename=path.name, data=path.read_bytes())
    mid = storage.create_meeting(pid, "Kickoff — Meridian card feed", "2026-01-12", t.to_text(), "vtt")
    pipe = Pipeline(storage, ReplayLLM())
    for stage in config.STAGES:
        pipe.run_stage(mid, stage)
        pipe.approve(mid, stage)
    wb = load_workbook(io.BytesIO(build_workbook(pipe, mid)))
    yield wb, pipe, mid
    storage.close()


def rows(ws):
    return list(ws.iter_rows(min_row=2, values_only=True))


def test_tabs_and_row_counts(exported):
    wb, _, _ = exported
    rec = load_recordings()[0].outputs
    assert wb.sheetnames == SHEET_NAMES
    assert len(rows(wb["Requirements"])) == len(rec["extract"]["requirements"])
    assert len(rows(wb["Decisions"])) == len(rec["extract"]["decisions"])
    assert len(rows(wb["Open Questions"])) == len(rec["extract"]["open_questions"])
    assert len(rows(wb["Epics"])) == len(rec["epics"]["epics"])
    assert len(rows(wb["Stories"])) == len(rec["stories"]["stories"])
    assert len(rows(wb["Tasks"])) == len(rec["tasks"]["tasks"])
    assert len(rows(wb["RAID"])) == len(rec["raid_email"]["raid"])


def test_formatting(exported):
    wb, _, _ = exported
    for name in SHEET_NAMES[1:-1]:
        ws = wb[name]
        assert ws.freeze_panes == "A2", name
        assert ws.auto_filter.ref.startswith("A1:"), name
        assert ws["A1"].font.bold and ws["A1"].font.name == "Arial", name


def test_dor_flags_and_missing_owners(exported):
    wb, pipe, mid = exported
    results = pipe.dor_results(mid)
    stories = wb["Stories"]
    for row in stories.iter_rows(min_row=2):
        sid, dor, issues = row[0].value, row[7].value, row[8].value
        assert (dor == "Pass") == results[sid].passed
        assert (row[0].fill.fgColor.rgb.endswith("FFF2CC")) == (dor != "Pass")
        if dor != "Pass":
            assert issues
    tasks = wb["Tasks"]
    owners = [r[4] for r in rows(tasks)]
    assert "TBD" in owners
    ac = stories["D2"].value
    assert ac.startswith("1. Given") and "[Negative]" in ac


def test_summary_formulas_point_at_the_right_columns(exported):
    """LibreOffice isn't available in CI, so check each formula's column by its header."""
    wb, _, _ = exported
    summary = {r[0]: r[1] for r in wb["Summary"].iter_rows(min_row=3, max_col=2, values_only=True) if r[0]}
    expected_headers = {"Story points": ("Stories", "Points"), "Tasks without an owner": ("Tasks", "Owner"),
                        "Definition of Ready pass rate": ("Stories", "DoR")}
    for label, (sheet, header) in expected_headers.items():
        formula = summary[label]
        col = re.search(rf"{sheet}!([A-Z]):", formula).group(1)
        assert wb[sheet][f"{col}1"].value == header, label
    for label, sheet in [("Requirements", "Requirements"), ("Open questions", "Open Questions"), ("Tasks", "Tasks")]:
        assert summary[label].startswith("=COUNTA(") and sheet in summary[label]
    assert summary["Model"] == "demo"


def test_email_tab(exported):
    wb, _, _ = exported
    lines = [r[0] for r in wb["Follow-up Email"].iter_rows(values_only=True)]
    assert lines[0].startswith("Subject: Recap")
    assert "Action items" in lines


def test_export_before_all_stages_are_done():
    storage = Storage(":memory:")
    pid = load_sample_project(storage)
    mid = storage.create_meeting(pid, "Empty", "2026-01-12", "Priya Nair: hi", "text")
    wb = load_workbook(io.BytesIO(build_workbook(Pipeline(storage), mid)))
    assert rows(wb["Stories"]) == []
    assert wb["Follow-up Email"]["A1"].value == "Not generated yet."


def test_filename():
    assert export_filename({"title": "Kickoff — Meridian / card feed!", "date": "2026-01-12"}) == \
        "2026-01-12_Kickoff-Meridian-card-feed.xlsx"
